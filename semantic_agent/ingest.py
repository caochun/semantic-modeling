from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from tempfile import TemporaryDirectory
import hashlib
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from .store import digest, now, tokens

SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".csv", ".xlsx", ".xls", ".md"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def decode_text(raw):
    encodings = ["utf-8-sig", "gb18030"]
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.insert(0, "utf-16")
    for encoding in encodings:
        try:
            return raw.decode(encoding)
        except UnicodeError:
            pass
    raise ValueError("无法识别文本编码")


def chunks(text, label, size=1800, overlap=3):
    """Keep complete rows/paragraphs where possible and give exact extracted line locators."""
    lines = text.splitlines()
    start = 0
    while start < len(lines):
        end, count = start, 0
        while end < len(lines) and (count < size or end == start):
            count += len(lines[end]) + 1
            end += 1
        body = "\n".join(lines[start:end]).strip()
        if body:
            # Very long table cells are split explicitly, never silently discarded.
            if len(body) > 7000:
                for offset in range(0, len(body), 5000):
                    yield f"{label}，提取行 {start + 1}-{end}，字符 {offset + 1}-{min(offset + 5000, len(body))}", body[offset:offset + 5000]
            else:
                yield f"{label}，提取行 {start + 1}-{end}", body
        if end >= len(lines):
            break
        start = max(start + 1, end - overlap)


def docx_text(path):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        body = root.find(W + "body")
        output = []
        for element in body if body is not None else []:
            if element.tag == W + "tbl":
                for row in element.findall(W + "tr"):
                    output.append(" | ".join("".join(n.text or "" for n in cell.iter(W + "t"))
                                             for cell in row.findall(W + "tc")))
            else:
                output.append("".join(n.text or "" for n in element.iter(W + "t")))
        return "\n".join(output)


def extract(path, soffice=""):
    suffix = path.suffix.lower()
    sections, notes = [], []
    if suffix == ".pdf":
        if shutil.which("pdftotext"):
            proc = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, timeout=90)
            if proc.returncode:
                raise ValueError("PDF 文本提取失败")
            pages = proc.stdout.decode("utf-8", errors="replace").split("\f")
            if pages and not pages[-1].strip():
                pages.pop()
        else:
            from pypdf import PdfReader
            pages = [p.extract_text() or "" for p in PdfReader(path).pages]
        empty = []
        for index, page in enumerate(pages, 1):
            if len(re.sub(r"\s", "", page)) < 25:
                empty.append(index)
            else:
                sections.append((f"PDF 第 {index} 页", page))
        if empty:
            notes.append("以下页面缺少可提取文本，图纸/扫描页需 OCR 或人工核读：" + ",".join(map(str, empty)))
    elif suffix == ".docx":
        sections.append(("DOCX 正文（表格以 | 分隔）", docx_text(path)))
        with zipfile.ZipFile(path) as archive:
            if any(name.startswith("word/media/") for name in archive.namelist()):
                notes.append("已提取正文与表格；内嵌图片未识别")
    elif suffix == ".doc":
        text = ""
        if shutil.which("textutil"):
            proc = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)], capture_output=True, timeout=90)
            if proc.returncode == 0:
                text = proc.stdout.decode("utf-8", errors="replace")
        executable = soffice or shutil.which("soffice")
        if not text.strip() and executable:
            with TemporaryDirectory(prefix="semantic-doc-") as tmp:
                profile = (Path(tmp) / "profile").as_uri()
                proc = subprocess.run([executable, f"-env:UserInstallation={profile}", "--headless",
                    "--convert-to", "txt:Text", "--outdir", tmp, str(path)], capture_output=True, timeout=120)
                out = Path(tmp) / (path.stem + ".txt")
                if out.exists():
                    text = out.read_text(encoding="utf-8-sig")
        if not text.strip():
            raise ValueError("未提取到 .doc 文本；可配置 SEMANTIC_SOFFICE_PATH 后重建索引")
        sections.append(("DOC 正文", text))
        notes.append("已提取文本；图片和复杂表格关系需核对原文件")
    elif suffix in {".xlsx", ".xls"}:
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    lines = [f"行{n}: " + " | ".join(str(v) if v is not None else "" for v in row)
                             for n, row in enumerate(sheet.iter_rows(values_only=True), 1)]
                    sections.append((f"工作表 {sheet.title}", "\n".join(lines)))
            finally:
                book.close()
        else:
            import xlrd
            book = xlrd.open_workbook(path)
            for sheet in book.sheets():
                lines = [f"行{n + 1}: " + " | ".join(map(str, sheet.row_values(n))) for n in range(sheet.nrows)]
                sections.append((f"工作表 {sheet.name}", "\n".join(lines)))
        notes.append("读取保存的单元格值；合并单元格、颜色和图形的含义需核对原表")
    else:
        sections.append(("文本", decode_text(path.read_bytes())))
    passages = [part for label, text in sections for part in chunks(text, label)]
    status = "partial" if notes and passages else "indexed" if passages else "needs_ocr"
    return passages, status, "；".join(notes)


def index_corpus(settings, store, progress=lambda *args: None):
    root = settings.doc_dir
    if not root.is_dir():
        raise ValueError("资料目录不存在")
    files = sorted(p for p in root.rglob("*") if p.is_file()
                   and p.suffix.lower() in SUPPORTED
                   and not any(part.startswith((".", "~$")) for part in p.relative_to(root).parts)
                   and p.resolve().is_relative_to(root.resolve()))
    if not files:
        raise ValueError("资料目录中没有支持的文件")
    seen, counts = set(), {"indexed": 0, "partial": 0, "needs_ocr": 0, "error": 0, "unchanged": 0}

    def prepare(path):
        relative = str(path.relative_to(root))
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        did = "doc_" + digest(relative + sha)
        prior = store.document(did)
        if prior and prior["status"] not in {"error", "needs_ocr"}:
            return path, relative, sha, did, None, prior["status"], prior["detail"]
        try:
            passages, status, detail = extract(path, settings.soffice)
        except Exception as exc:
            passages, status, detail = [], "error", str(exc)[:250]
        return path, relative, sha, did, passages, status, detail

    with ThreadPoolExecutor(max_workers=4) as pool:
        for index, result in enumerate(pool.map(prepare, files), 1):
            path, relative, sha, did, passages, status, detail = result
            seen.add(did)
            stat = path.stat()
            with store.connect() as db:
                db.execute("UPDATE documents SET active=0 WHERE path=? AND id!=?", (relative, did))
                db.execute("INSERT OR REPLACE INTO documents VALUES(?,?,?,?,?,?,?,?,?)", (
                    did, relative, sha, stat.st_size, stat.st_mtime_ns, 1, status, detail, now()))
                if passages is not None:
                    db.execute("DELETE FROM passage_fts WHERE passage_id IN (SELECT id FROM passages WHERE document_id=?)", (did,))
                    db.execute("DELETE FROM passages WHERE document_id=?", (did,))
                    for locator, text in passages:
                        pid = "p_" + digest(did + locator + text)
                        db.execute("INSERT INTO passages VALUES(?,?,?,?)", (pid, did, locator, text))
                        db.execute("INSERT INTO passage_fts VALUES(?,?,?)", (pid,
                            " ".join(tokens(relative)), " ".join(tokens(text))))
                else:
                    counts["unchanged"] += 1
            counts[status] += 1
            if index % 10 == 0 or index == len(files):
                progress(index, len(files), relative)
    with store.connect() as db:
        active = db.execute("SELECT id FROM documents WHERE active=1").fetchall()
        for row in active:
            if row["id"] not in seen:
                db.execute("UPDATE documents SET active=0 WHERE id=?", (row["id"],))
    return {"files": len(files), "counts": counts, **store.stats()}
